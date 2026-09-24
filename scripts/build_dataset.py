# -*- coding: utf-8 -*-
"""生データ3種を、ブラウザに同梱できる1つのGeoJSONに落とす。

  地域危険度SHP (EPSG:2451, 5192ポリゴン)
  + 地域危険度CSV (地盤分類。SHPには入っていない)
  + 東京消防庁 出火危険度XLSX (冬夕 / 夏昼)
  -> public/data/chochome.json  (WGS84)

位置情報を端末から出さない設計なので、判定に必要なものは全てクライアントに置く。
"""
import io
import json
import os
import sys
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
import shapefile
from norm import normalize, variants
from pyproj import Transformer
from shapely.geometry import shape as to_shape
from shapely.ops import transform as shp_transform

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
RAW = os.path.join(ROOT, "data", "raw")
OUT_DIR = os.path.join(ROOT, "public", "data")

# 平面直角座標系第9系での簡略化許容誤差(m)。町丁目単位の相対評価なので数mのズレは無害。
TOLERANCE_M = 5.0
COORD_DIGITS = 6  # 約0.1m。これ以上は無駄

# EPSG:2451 = JGD2000 / Japan Plane Rectangular CS IX（.prjのCentral_Meridian 139.8333と一致）
TF = Transformer.from_crs("EPSG:2451", "EPSG:4326", always_xy=True)


def load_jiban():
    """CSVから地盤分類を拾う。キーは 区市町名+町丁目名（同一出典なのでSHPと一致するはず）。"""
    lines = open(os.path.join(RAW, "kikendo9.csv"), "rb").read().decode("cp932").splitlines()
    out = {}
    for l in lines[1:]:
        c = l.split(",")
        out[normalize(c[0] + c[1])] = c[2]
    return out


def load_shukka():
    """出火危険度。ヘッダ3行、4行目からデータ。列は inspect_raw.py の出力で確認済み。

    [0] 町丁目名 / [1..5] 冬夕 要因別 / [6] 冬夕 木造 / [7] 冬夕 非木造
    [8] 冬夕 総合 / [9] 夏昼 総合
    """
    wb = openpyxl.load_workbook(os.path.join(RAW, "tfd_shukka_choChome.xlsx"), read_only=True)
    ws = wb["Sheet1"]
    out = {}

    def num(v):
        try:
            return int(str(v).strip())
        except (TypeError, ValueError):
            return None

    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        out[normalize(str(row[0]))] = {
            "shukkaWinterEve": num(row[8]),
            "shukkaSummerNoon": num(row[9]),
            "shukkaWoodWinterEve": num(row[6]),
            "shukkaNonWoodWinterEve": num(row[7]),
        }
    return out


def lookup(table, key):
    for v in variants(key):
        if v in table:
            return table[v]
    return None


def main():
    jiban = load_jiban()
    shukka = load_shukka()

    shp_path = [os.path.join(RAW, "kikendo9_shp", f)
                for f in os.listdir(os.path.join(RAW, "kikendo9_shp"))
                if f.endswith(".shp")][0]
    sf = shapefile.Reader(shp_path, encoding="cp932")

    # 災害時活動困難係数だけ元データに順位が無いので、ここで作る。
    # 他の3指標と同じ「1 = 都内で最も厳しい」向きに揃える。
    difficulty_order = {}
    for pos, rec in enumerate(sorted(sf.records(), key=lambda r: -r["災害_係"]), start=1):
        difficulty_order[rec["ID"]] = pos

    features = []
    stats = Counter()
    pts_before = pts_after = 0

    for rec, shp in zip(sf.records(), sf.shapes()):
        city = rec["区市町村名"].strip()
        town = rec["町丁目名"].strip()
        key = normalize(city + town)

        geom = to_shape(shp.__geo_interface__)
        pts_before += len(shp.points)
        if not geom.is_valid:
            geom = geom.buffer(0)
            stats["ジオメトリ修復"] += 1
        geom = geom.simplify(TOLERANCE_M, preserve_topology=True)
        geom = shp_transform(lambda x, y, z=None: TF.transform(x, y), geom)
        gj = json.loads(json.dumps(geom.__geo_interface__))
        gj = round_coords(gj)
        pts_after += count_points(gj)

        j = jiban.get(key)
        if j is None:
            stats["地盤分類が引けない"] += 1
        s = lookup(shukka, key)
        if s is None:
            stats["出火危険度が引けない"] += 1
            s = {"shukkaWinterEve": None, "shukkaSummerNoon": None,
                 "shukkaWoodWinterEve": None, "shukkaNonWoodWinterEve": None}

        features.append({
            "type": "Feature",
            "geometry": gj,
            "properties": {
                "id": rec["ID"],
                "city": city,
                "town": town,
                "ground": j,
                # 危険量は連続値。5段階ランクではなくこちらでスコアリングする。
                # 順位(1 = 都内で最も危険)は、★を「5,192町丁目の中での相対値」として
                # 引き直すのにそのまま使える。
                "collapseQty": round(rec["建物_危"], 4),
                "collapseOrder": rec["建物_順"],
                "collapseRank": rec["建物_ラ"],
                "fireQty": round(rec["火災_危"], 4),
                "fireOrder": rec["火災_順"],
                "fireRank": rec["火災_ラ"],
                "difficulty": round(rec["災害_係"], 4),
                "difficultyOrder": difficulty_order[rec["ID"]],
                "totalQty": round(rec["総合_危"], 4),
                "totalOrder": rec["総合_順"],
                "totalRank": rec["総合_ラ"],
                **s,
            },
        })
        stats["町丁目"] += 1

    os.makedirs(OUT_DIR, exist_ok=True)
    # 拡張子は .geojson ではなく .json。application/geo+json はCDNの圧縮対象
    # MIMEタイプに入っていないことがあり、4.65MBが無圧縮で配信される恐れがある。
    out_path = os.path.join(OUT_DIR, "chochome.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump({"type": "FeatureCollection", "features": features},
                  f, ensure_ascii=False, separators=(",", ":"))

    print("=== 出力 ===")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    print(f"  頂点数: {pts_before:,} -> {pts_after:,} "
          f"({pts_after / pts_before * 100:.1f}%, tolerance={TOLERANCE_M}m)")
    size = os.path.getsize(out_path)
    print(f"  {out_path}: {size/1024/1024:.2f} MB")
    import gzip
    gz = len(gzip.compress(open(out_path, "rb").read(), 6))
    print(f"  gzip後: {gz/1024/1024:.2f} MB  <- 実際に転送される量")


def round_coords(g):
    def r(c):
        if isinstance(c[0], (int, float)):
            return [round(c[0], COORD_DIGITS), round(c[1], COORD_DIGITS)]
        return [r(x) for x in c]
    g["coordinates"] = r(g["coordinates"])
    return g


def count_points(g):
    def c(x):
        if isinstance(x[0], (int, float)):
            return 1
        return sum(c(i) for i in x)
    return c(g["coordinates"])


if __name__ == "__main__":
    main()
