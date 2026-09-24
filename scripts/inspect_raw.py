# -*- coding: utf-8 -*-
"""生データの中身を実測する。推測せず、必ずここの出力を根拠にする。"""
import glob
import io
import json
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

RAW = os.path.join(os.path.dirname(__file__), "..", "data", "raw")


def head(title):
    print("\n" + "=" * 70)
    print(title)
    print("=" * 70)


# ---------------------------------------------------------------- CSV
head("地域危険度（第9回）CSV")
csv_bytes = open(os.path.join(RAW, "kikendo9.csv"), "rb").read()
text = csv_bytes.decode("cp932")
lines = text.splitlines()
cols = lines[0].split(",")
print("行数(ヘッダ除く):", len(lines) - 1)
print("カラム数:", len(cols))
for i, c in enumerate(cols):
    print(f"  [{i}] {c}")
print("\nサンプル3行:")
for l in lines[1:4]:
    print("  ", l)

# 地盤分類の値域
gb = {}
for l in lines[1:]:
    v = l.split(",")[2]
    gb[v] = gb.get(v, 0) + 1
print("\n地盤分類の値域:", json.dumps(gb, ensure_ascii=False))

# ---------------------------------------------------------------- SHP
head("地域危険度（第9回）SHP")
import shapefile  # noqa: E402

shp_path = glob.glob(os.path.join(RAW, "kikendo9_shp", "*.shp"))[0]
print("ファイル:", os.path.basename(shp_path))
prj = open(shp_path[:-4] + ".prj", encoding="utf-8", errors="replace").read()
print("prj:", prj[:400])

sf = shapefile.Reader(shp_path, encoding="cp932")
print("フィーチャ数:", len(sf))
print("shapeType:", sf.shapeType)
print("bbox:", sf.bbox)
print("フィールド:")
for f in sf.fields[1:]:
    print("  ", f)
rec = sf.record(0)
print("\n先頭レコード:", json.dumps(dict(rec.as_dict()), ensure_ascii=False, default=str))
s = sf.shape(0)
print("先頭ジオメトリ: parts=", len(s.parts), " points=", len(s.points))
print("先頭座標3点:", s.points[:3])

# 頂点総数（クライアント同梱サイズの見積もりに使う）
total_pts = 0
for s in sf.iterShapes():
    total_pts += len(s.points)
print("全頂点数:", total_pts)

# ---------------------------------------------------------------- XLSX
head("東京消防庁 町丁目別出火危険度 XLSX")
import openpyxl  # noqa: E402

wb = openpyxl.load_workbook(os.path.join(RAW, "tfd_shukka_choChome.xlsx"), read_only=True)
print("シート:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n--- シート '{name}' ({ws.max_row} 行 x {ws.max_column} 列) ---")
    for r, row in enumerate(ws.iter_rows(max_row=12, values_only=True)):
        print(f"  [{r}]", row)
